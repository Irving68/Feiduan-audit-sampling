from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from audit_sampling.agent_protocol import normalize_agent_payload, validate_agent_fields
from audit_sampling.excel_export import export_excel
from audit_sampling.matcher import match_fields
from audit_sampling.preview import render_preview


class CoreFlowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.manifest = {
            "job": {"job_id": "abc12345", "mode": "non-whole"},
            "files": [{"file_id": "file_001", "name": "sample.png"}],
            "pages": [{"page_id": "file_001_p001", "file_id": "file_001", "filename": "sample.png", "source_page": 1, "image": "pages/file_001_p001.jpg", "width": 1000, "height": 1400}],
            "ocr": [{"page_id": "file_001_p001", "file_id": "file_001", "filename": "sample.png", "source_page": 1, "words": [
                {"word_id": "w1", "text": "付款人户名：测试有很公司", "confidence": .99, "bbox": {"left": 10, "top": 20, "width": 300, "height": 40}},
                {"word_id": "w2", "text": "小写：870,339.28元", "confidence": .99, "bbox": {"left": 10, "top": 80, "width": 220, "height": 40}},
                {"word_id": "w3", "text": "合计：870,339.28元", "confidence": .93, "bbox": {"left": 10, "top": 140, "width": 220, "height": 40}},
            ]}],
            "fields": [], "artifacts": {}, "summary": {},
        }

    def test_agent_group_and_match(self) -> None:
        fields = normalize_agent_payload({"fields": [
            {"document_type": "银行回单", "file": "sample.png", "file_id": "file_001", "group_id": "g1", "name": "付款人户名", "value": "测试有限公司", "raw_value": "测试有很公司", "source_text": "付款人户名：测试有很公司", "page_hint": 1, "evidence_word_ids": ["w1"], "confidence": .96, "importance": "high"},
            {"document_type": "银行回单", "file": "sample.png", "file_id": "file_001", "group_id": "g1", "name": "金额", "value": "870339.28", "raw_value": "870,339.28", "source_text": "小写：870,339.28元", "page_hint": 1, "evidence_word_ids": ["w2"], "confidence": .98, "importance": "high"},
        ]})
        matched = match_fields(self.manifest, fields)
        self.assertTrue(all(item["status"] == "confirmed" for item in matched))
        self.assertTrue(all(item["page_id"] == "file_001_p001" for item in matched))

    def test_corrected_value_uses_raw_ocr_for_highlight(self) -> None:
        fields = normalize_agent_payload({"fields": [
            {"document_type": "银行回单", "file": "sample.png", "group_id": "g1", "name": "付款人户名",
             "value": "测试有限公司", "raw_value": "测试有很公司", "source_text": "付款人户名：测试有很公司",
             "page_hint": 1, "confidence": .96, "importance": "high", "correction": "修正OCR错字"}
        ]})
        matched = match_fields(self.manifest, fields)
        self.assertEqual(matched[0]["value"], "测试有限公司")
        self.assertEqual(matched[0]["raw_value"], "测试有很公司")
        self.assertEqual(matched[0]["evidence_source"], "raw_value")
        self.assertEqual(matched[0]["status"], "confirmed")

    def test_evidence_word_id_resolves_duplicate_value(self) -> None:
        fields = normalize_agent_payload({"fields": [
            {"document_type": "银行回单", "file": "sample.png", "file_id": "file_001", "group_id": "g1", "name": "合计金额",
             "value": "870339.28", "raw_value": "870,339.28", "source_text": "合计：870,339.28元",
             "page_hint": 1, "evidence_word_ids": ["w3"], "confidence": .98, "importance": "high"}
        ]})
        self.assertEqual(validate_agent_fields(self.manifest, fields), [])
        matched = match_fields(self.manifest, fields)
        self.assertEqual(matched[0]["evidence_word_ids"], ["w3"])
        self.assertEqual(matched[0]["bbox"]["top"], 140)
        self.assertEqual(matched[0]["status"], "confirmed")
        self.assertEqual(matched[0]["confidence"], .93)

    def test_validation_requires_evidence_ids(self) -> None:
        fields = normalize_agent_payload({"fields": [
            {"document_type": "银行回单", "file": "sample.png", "file_id": "file_001", "group_id": "g1", "name": "金额",
             "value": "870339.28", "raw_value": "870,339.28", "source_text": "小写：870,339.28元",
             "page_hint": 1, "confidence": .98, "importance": "high"}
        ]})
        errors = validate_agent_fields(self.manifest, fields)
        self.assertTrue(any("evidence_word_ids" in error for error in errors))

    def test_missing_value_keeps_unique_page_hint(self) -> None:
        fields = normalize_agent_payload({"fields": [
            {"document_type": "银行回单", "file": "sample.png", "group_id": "g1", "name": "备注",
             "value": "OCR中不存在的内容", "page_hint": 1, "confidence": .7, "importance": "high"}
        ]})
        matched = match_fields(self.manifest, fields)
        self.assertEqual(matched[0]["status"], "missing")
        self.assertEqual(matched[0]["page_id"], "file_001_p001")
        self.assertEqual(matched[0]["bboxes"], [])

    def test_excel_and_preview(self) -> None:
        self.manifest["fields"] = match_fields(self.manifest, normalize_agent_payload({"fields": [
            {"document_type": "银行回单", "file": "sample.png", "file_id": "file_001", "group_id": "g1", "name": "金额", "value": "870339.28", "raw_value": "870,339.28", "source_text": "小写：870,339.28元", "page_hint": 1, "evidence_word_ids": ["w2"], "confidence": .88, "importance": "high"}
        ]}))
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); (root / "pages").mkdir(); (root / "pages/file_001_p001.jpg").write_bytes(b"test")
            excel = export_excel(self.manifest, root / "result.xlsx")
            wb = openpyxl.load_workbook(excel); ws = wb["多套资料审计"]
            self.assertEqual(ws["A1"].value, "资料类型")
            self.assertEqual(ws["A3"].value, "字段值\nsample.png")
            self.assertEqual(ws["B4"].value, "88%")
            self.assertEqual(ws["B4"].fill.fill_type, "solid")
            self.assertEqual(ws["B3"].hyperlink.target, "links/field-field_0001.html")
            preview = render_preview(self.manifest, root)
            page = preview.read_text(encoding="utf-8")
            self.assertIn("王二审计抽凭 v1.2.0-beta", page)
            self.assertIn("需人工复核", page)
            self.assertIn("file_001_p001", page)
            self.assertIn("展开全部", page)
            self.assertIn("适应宽度", page)
            self.assertTrue((root / "open-preview.bat").exists())
            self.assertTrue((root / "links/field-field_0001.html").exists())
            self.assertIn("preview.html?field=field_0001", (root / "links/field-field_0001.html").read_text(encoding="utf-8"))

    def test_preview_groups_multiple_files_and_instances(self) -> None:
        manifest = {
            "job": {"job_id": "multi001", "mode": "whole-set"},
            "files": [
                {"file_id": "file_001", "name": "sample-a.pdf"},
                {"file_id": "file_002", "name": "sample-b.pdf"},
            ],
            "pages": [
                {"page_id": "file_001_p001", "filename": "sample-a.pdf", "source_page": 1, "image": "pages/a.jpg", "width": 100, "height": 100},
                {"page_id": "file_002_p001", "filename": "sample-b.pdf", "source_page": 1, "image": "pages/b.jpg", "width": 100, "height": 100},
            ],
            "fields": [
                {"file": "sample-a.pdf", "document_type": "发票", "group_id": "a-1", "name": "金额", "value": "100", "page": 1, "page_id": "file_001_p001", "confidence": .98, "status": "confirmed", "bboxes": []},
                {"file": "sample-a.pdf", "document_type": "发票", "group_id": "a-2", "name": "金额", "value": "200", "page": 1, "page_id": "file_001_p001", "confidence": .98, "status": "confirmed", "bboxes": []},
                {"file": "sample-b.pdf", "document_type": "银行回单", "group_id": "b-1", "name": "付款方", "value": "测试公司", "page": 1, "page_id": "file_002_p001", "confidence": .88, "status": "confirmed", "bboxes": []},
            ],
        }
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            preview = render_preview(manifest, root)
            page = preview.read_text(encoding="utf-8")
            self.assertIn("sample-a.pdf", page)
            self.assertIn("sample-b.pdf", page)
            self.assertIn("group_id", page)
            self.assertIn("上一页", page)
            self.assertIn("下一页", page)
            self.assertIn("selectLinkTarget", page)
            self.assertIn("const isOpen=g.fields.some", page)
            self.assertIn("showPage(0,false);selectLinkTarget()", page)
            self.assertNotIn("showPage(0,true)", page)

    def test_excel_aligns_fields_by_name_and_adds_comment(self) -> None:
        manifest = {
            "job": {"job_id": "align001", "mode": "non-whole"},
            "files": [{"file_id": "file_001", "name": "sample.pdf"}],
            "fields": [
                {"file": "sample.pdf", "file_id": "file_001", "document_type": "发票", "group_id": "g1", "name": "日期", "value": "2026-01-01", "field_id": "field_0001", "page_id": "p1", "confidence": .98, "review_reasons": []},
                {"file": "sample.pdf", "file_id": "file_001", "document_type": "发票", "group_id": "g1", "name": "金额", "value": "100", "raw_value": "1OO", "correction": "修正OCR字母", "field_id": "field_0002", "page_id": "p1", "confidence": .95, "review_reasons": []},
                {"file": "sample.pdf", "file_id": "file_001", "document_type": "发票", "group_id": "g2", "name": "金额", "value": "200", "field_id": "field_0003", "page_id": "p2", "confidence": .92, "review_reasons": []},
            ],
        }
        with tempfile.TemporaryDirectory() as folder:
            excel = export_excel(manifest, Path(folder) / "result.xlsx")
            wb = openpyxl.load_workbook(excel); ws = wb["多套资料审计"]
            self.assertEqual(ws["B4"].value, "-")
            self.assertEqual(ws["C4"].value, "200")
            self.assertIn("来源OCR：1OO", ws["C3"].comment.text)
            self.assertEqual(ws["C3"].hyperlink.target, "links/field-field_0002.html")


if __name__ == "__main__":
    unittest.main()
