"""王二审计抽凭 v1.2.0-beta 的 Excel 输出适配层。

保留既有标签、合并单元格和准确度行，只增强字段对齐、证据链接和复核追溯。
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any
from urllib.parse import quote

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REVIEW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
LINK_COLOR = "0563C1"


def safe_cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def _field_key(value: Any) -> str:
    return re.sub(r"[\s，,。．.：:；;（）()【】\[\]_\-]+", "", str(value or "")).lower()


def groups_from_fields(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    """将扁平字段按文件、资料实例重建，并保留完整证据元数据。"""

    grouped: dict[tuple[str, str, str], dict[str, Any]] = {}
    for field in manifest.get("fields", []):
        file_key = str(field.get("file_id") or field.get("file") or "")
        key = (
            file_key,
            str(field.get("document_type", "未知资料") or "未知资料"),
            str(field.get("group_id", "") or "default"),
        )
        group = grouped.setdefault(key, {
            "file_id": str(field.get("file_id", "") or ""),
            "filename": str(field.get("file", "") or ""),
            "doc_type": key[1],
            "group_id": key[2],
            "page": field.get("page"),
            "audit_key_fields": [],
        })
        group["audit_key_fields"].append(dict(field))

    files = manifest.get("files", [])
    name_counts = Counter(str(item.get("name", "") or "") for item in files)
    name_seen: Counter[str] = Counter()
    records = []
    used_group_ids = set()
    for item in files:
        file_id = str(item.get("file_id", "") or "")
        filename = str(item.get("name", "") or "")
        name_seen[filename] += 1
        display = filename if name_counts[filename] <= 1 else f"{filename}（{name_seen[filename]}）"
        groups = [group for group in grouped.values() if group["file_id"] == file_id]
        if not groups:
            groups = [group for group in grouped.values() if not group["file_id"] and group["filename"] == filename]
        if groups:
            for group in groups:
                group["display_filename"] = display
                used_group_ids.add(id(group))
            records.append({"file_id": file_id, "filename": filename, "display_filename": display, "document_groups": groups})

    for group in grouped.values():
        if id(group) in used_group_ids:
            continue
        group["display_filename"] = group["filename"]
        records.append({
            "file_id": group["file_id"],
            "filename": group["filename"],
            "display_filename": group["filename"],
            "document_groups": [group],
        })
    return records


def _style_label(cell: Any) -> None:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _link_filename(fragment: str) -> str:
    return f"{fragment.replace('=', '-', 1)}.html"


def _set_link(cell: Any, fragment: str, link_base: str = "links") -> None:
    cell.hyperlink = f"{link_base}/{_link_filename(fragment)}"
    cell.font = Font(color=LINK_COLOR, underline="single", bold=cell.font.bold)


def _field_map(group: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {_field_key(field.get("name")): field for field in group.get("audit_key_fields", [])}


def _field_names(groups: list[dict[str, Any]]) -> list[str]:
    names: list[str] = []
    seen = set()
    for group in groups:
        for field in group.get("audit_key_fields", []):
            key = _field_key(field.get("name"))
            if key and key not in seen:
                seen.add(key)
                names.append(str(field.get("name", "") or ""))
    return names


def _needs_review(field: dict[str, Any] | None) -> bool:
    if field is None:
        return True
    return bool(field.get("review_reasons")) or float(field.get("confidence", 0) or 0) < .9


def _accuracy_text(field: dict[str, Any] | None) -> str:
    if field is None:
        return "--（字段缺失）"
    percent = f"{int(round(float(field.get('confidence', 0) or 0) * 100))}%"
    reason = str(field.get("review_reason", "") or "")
    return f"{percent}（{reason}）" if reason else percent


def _comment_text(field: dict[str, Any]) -> str:
    lines = []
    raw_value = str(field.get("raw_value", "") or field.get("source_text", "") or "")
    if raw_value:
        lines.append(f"来源OCR：{raw_value}")
    if field.get("correction"):
        lines.append(f"Agent修正：{field.get('value', '')}")
        lines.append(f"修正说明：{field.get('correction', '')}")
    reasons = field.get("review_reasons", []) or []
    if reasons:
        lines.append(f"复核原因：{'、'.join(str(reason) for reason in reasons)}")
    return "\n".join(lines)


def _write_value(cell: Any, field: dict[str, Any] | None, link_base: str = "preview.html") -> None:
    if field is None:
        cell.value = "-"
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        return
    cell.value = safe_cell_value(field.get("value", ""))
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if field.get("field_id") and field.get("page_id"):
        _set_link(cell, f"field={quote(str(field['field_id']))}", link_base)
    if field.get("correction") or _needs_review(field):
        note = _comment_text(field)
        if note:
            cell.comment = Comment(note, "WangEr Audit Sampling")


def _write_accuracy(ws: Any, row: int, groups: list[dict[str, Any]], names: list[str], start_col: int) -> None:
    ws.cell(row, 1).value = "识别准确度"
    _style_label(ws.cell(row, 1))
    maps = [_field_map(group) for group in groups]
    for offset, name in enumerate(names):
        fields = [mapping.get(_field_key(name)) for mapping in maps]
        cell = ws.cell(row, start_col + offset)
        cell.value = " / ".join(_accuracy_text(field) for field in fields)
        _style_label(cell)
        if any(_needs_review(field) for field in fields):
            cell.fill = REVIEW_FILL


def export_excel(manifest: dict[str, Any], output_path: str | Path) -> Path:
    mode = manifest["job"].get("mode")
    records = groups_from_fields(manifest)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "多套资料审计"
    current_row = 1
    link_base = "links"
    preview_file = manifest.get("artifacts", {}).get("preview_file")
    if sys.platform == "darwin" and preview_file:
        link_base = (Path(preview_file).resolve().parent / "links").as_uri()

    if mode == "whole-set":
        for result_index, result in enumerate(records):
            if result_index:
                current_row += 1
            by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
            for group in result["document_groups"]:
                by_type[group.get("doc_type", "Unknown")].append(group)
            max_instances = max((len(value) for value in by_type.values()), default=1)
            label = ws.cell(current_row, 1)
            label.value = f"资料类型\n{result['display_filename']}"
            _style_label(label)
            if result.get("file_id"):
                _set_link(label, f"file={quote(str(result['file_id']))}", link_base)

            names_by_type = {doc_type: _field_names(groups) for doc_type, groups in by_type.items()}
            col = 2
            for doc_type in by_type:
                names = names_by_type[doc_type]
                if names:
                    ws.merge_cells(start_row=current_row, start_column=col, end_row=current_row, end_column=col + len(names) - 1)
                    ws.cell(current_row, col).value = doc_type
                    _style_label(ws.cell(current_row, col))
                    col += len(names)
                col += 1

            ws.cell(current_row + 1, 1).value = "资料字段"
            _style_label(ws.cell(current_row + 1, 1))
            col = 2
            for doc_type in by_type:
                for name in names_by_type[doc_type]:
                    cell = ws.cell(current_row + 1, col)
                    cell.value = name
                    _style_label(cell)
                    col += 1
                col += 1

            for instance in range(max_instances):
                row = current_row + 2 + instance
                ws.cell(row, 1).value = "字段值" if instance == 0 else ""
                if instance == 0:
                    _style_label(ws.cell(row, 1))
                col = 2
                for doc_type, groups in by_type.items():
                    mapping = _field_map(groups[instance]) if instance < len(groups) else {}
                    for name in names_by_type[doc_type]:
                        _write_value(ws.cell(row, col), mapping.get(_field_key(name)), link_base)
                        col += 1
                    col += 1

            accuracy_row = current_row + 2 + max_instances
            col = 2
            for doc_type, groups in by_type.items():
                names = names_by_type[doc_type]
                _write_accuracy(ws, accuracy_row, groups, names, col)
                col += len(names) + 1
            current_row = accuracy_row + 1
    else:
        by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for result in records:
            for group in result["document_groups"]:
                by_type[group.get("doc_type", "Unknown")].append(group)
        for type_index, (doc_type, groups) in enumerate(by_type.items()):
            if type_index:
                current_row += 1
            names = _field_names(groups)
            ws.cell(current_row, 1).value = "资料类型"
            _style_label(ws.cell(current_row, 1))
            if names:
                ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=1 + len(names))
                ws.cell(current_row, 2).value = doc_type
                _style_label(ws.cell(current_row, 2))
            ws.cell(current_row + 1, 1).value = "资料字段"
            _style_label(ws.cell(current_row + 1, 1))
            for offset, name in enumerate(names, start=2):
                cell = ws.cell(current_row + 1, offset)
                cell.value = name
                _style_label(cell)

            for instance, group in enumerate(groups):
                row = current_row + 2 + instance
                label = ws.cell(row, 1)
                label.value = f"字段值\n{group.get('display_filename', group.get('filename', ''))}"
                _style_label(label)
                first_field = next(iter(group.get("audit_key_fields", [])), None)
                if first_field and first_field.get("field_id"):
                    _set_link(label, f"field={quote(str(first_field['field_id']))}", link_base)
                mapping = _field_map(group)
                for offset, name in enumerate(names, start=2):
                    _write_value(ws.cell(row, offset), mapping.get(_field_key(name)), link_base)
            accuracy_row = current_row + 2 + len(groups)
            _write_accuracy(ws, accuracy_row, groups, names, 2)
            current_row = accuracy_row + 1

    ws.column_dimensions["A"].width = 20
    for col in range(2, 50):
        ws.column_dimensions[get_column_letter(col)].width = 16
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target
